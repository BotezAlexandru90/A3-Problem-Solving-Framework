from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .db import CHILD_TABLES
from .model import STEPS

CHILD_TITLES = {
    "is_is_not": "Is / Is Not",
    "hypothesis": "Cause hypotheses and verification",
    "why_step": "Why-Why chain",
    "action": "Action plan",
    "measurement": "KPI measurements",
    "read_across": "Read Across",
    "lesson": "Lessons Learned",
}


def _md_cell(v):
    return str(v if v is not None else "").replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def to_markdown(bundle, evaluation):
    a3 = bundle["a3"]
    out = [f"# A3 #{a3['id']}: {a3['title']}", "",
           f"Status: {a3['status']} | Last update: {a3['updated_at']} by {a3['updated_by'] or '-'}", ""]
    for step, label, fields in STEPS:
        out += [f"## {label}", ""]
        for name, flabel, _ in fields:
            value = (a3.get(name) or "").strip() or "_not yet provided_"
            out += [f"**{flabel}:**", "", value, ""]
        gate = evaluation[step]
        out.append(f"_Quality gate: {'passed' if gate['ready'] else 'open'}_")
        for c in gate["checks"]:
            if not c.passed:
                out.append(f"- [ ] {c.label}" + (" (advice)" if c.level == "advice" else ""))
        out.append("")
    for table, title in CHILD_TITLES.items():
        rows = bundle[table]
        out += [f"## {title}", ""]
        if not rows:
            out += ["_none_", ""]
            continue
        cols = ["id"] + CHILD_TABLES[table]
        out.append("| " + " | ".join(cols) + " |")
        out.append("|" + "---|" * len(cols))
        for r in rows:
            out.append("| " + " | ".join(_md_cell(r[c]) for c in cols) + " |")
        out.append("")
    return "\n".join(out)


def _xl(v):
    s = "" if v is None else str(v)
    if s[:1] in ("=", "+", "-", "@"):
        s = "'" + s  # prevent formula injection
    return s[:32767]


def to_xlsx(bundle, evaluation):
    wb = Workbook()
    ws = wb.active
    ws.title = "A3"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="FF5F00")
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 100
    a3 = bundle["a3"]
    ws.append(["A3 #", a3["id"]])
    ws.append(["Status", a3["status"]])
    ws.append(["Last update", f"{a3['updated_at']} by {a3['updated_by']}"])
    for step, label, fields in STEPS:
        ws.append([])
        ws.append([label, "Gate: " + ("passed" if evaluation[step]["ready"] else "open")])
        for c in ws[ws.max_row]:
            c.font, c.fill = head, fill
        for name, flabel, _ in fields:
            ws.append([flabel, _xl(a3.get(name))])
            for c in ws[ws.max_row]:
                c.alignment = wrap

    for table, title in CHILD_TITLES.items():
        sh = wb.create_sheet(title[:31].replace("/", "-"))
        cols = ["id"] + CHILD_TABLES[table]
        sh.append(cols)
        for c in sh[1]:
            c.font, c.fill = head, fill
        for r in bundle[table]:
            sh.append([_xl(r[c]) for c in cols])
        for i in range(1, len(cols) + 1):
            sh.column_dimensions[sh.cell(1, i).column_letter].width = 22

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
